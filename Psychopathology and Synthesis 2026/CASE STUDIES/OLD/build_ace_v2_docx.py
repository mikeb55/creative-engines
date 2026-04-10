# -*- coding: utf-8 -*-
"""Generate V2 ACE case Word documents (polished narratives; do not overwrite ACE-Case*.md)."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_DIR = Path(__file__).resolve().parent

CASES = [
    {
        "filename": "V2-Leah-Shame.docx",
        "title": "ACE-Case Leah — Shame",
        "paras": [
            "Leah is twenty-three and self-referred to a low-cost counselling service. She lives with parents who tell professionals she is capable but cannot “get going.” In their world, adulthood is tied to steady employment and self-sufficiency; her repeated job losses sit inside a family story that she is “wasting her education,” which deepens shame and hiding. Over roughly two years she has taken jobs in shops and cafés and left each within weeks. Busy floors, unpredictable shifts, and managerial scrutiny flood her with anxiety; she calls in sick, then stops going. Each exit sharpens a private sense of failure. She retreats to her room much of the day, sleeps at odd hours—often awake late on her phone, asleep past morning commitments—and has let friendships lapse, saying others should not have to “carry” her.",
            "She describes low mood, little energy, and intrusive thoughts that she ruins everything and that effort is pointless. After social contact she ruminates on whether she said the wrong thing. A friend has wondered aloud about autism or ADHD. Leah recognises sensory overload and exhaustion cycles but refuses formal assessment, fearing a label would define her. Any exploration of neurodivergence-relevant themes needs to stay tentative, descriptive, and co-created—not diagnostic or labelling.",
            "In the opening appointment, when risk is explored calmly, she denies past self-harm yet admits passive thoughts of not wanting to be alive. She mentions a recent row with her parents after which she stood at the household medicine cupboard wanting “everything to stop,” then stepped away within minutes without taking anything. She has not disclosed this to family or GP. She is unsure therapy can help—part of her hopes to “get back on track,” part expects to drop out as she has dropped jobs.",
            "In the room she can seem muted yet thoughtful, with brief moments of warmth or humour. She has already exercised agency by self-referring and by tolerating direct risk questions. The clinical question for trainees is how much structure she can bear early on, how explicitly to hold risk without shaming her, and where neurodevelopmental curiosity belongs alongside mood, family dynamics, and avoidance.",
        ],
        "trainee": "In this exercise, please do not use teaching time asking about missing background details; instead, agree as a group to “fill in the gaps” so you can focus on completing the clinical task.",
        "tasks": [
            "Agree initial assessment priorities and immediate safety steps for the first two sessions.",
            "Draft a brief working formulation and a plausible focus for sessions one through six to eight, stating what realistic change would look like.",
            "Identify one respectful way to explore neurodivergence-relevant questions without bypassing shame or active risk.",
        ],
        "ace": "The vignette chains observable behaviour, family and cultural framing of independence, shame, and a specific low-lethality but serious risk moment so that reasoning about assessment and early pacing is explicit; tasks force formulation, sequencing, and integration of ND-sensitive enquiry with safeguarding, supporting strong marks on analytical clarity, clinical completeness, context & ethics, and educational value.",
    },
    {
        "filename": "V2-Ousmane-Overload.docx",
        "title": "ACE-Case Ousmane — Overload",
        "paras": [
            "Ousmane is thirty-six and raising two children, aged six and nine, alone. A local support service has referred him after repeated no-shows at appointments and rising rent arrears. The school has flagged patchy attendance, homework not returned, and concern that the home environment may not be steady enough for learning. Teachers describe mornings when the children arrive tired, without lunch money or signed forms. Ousmane is not hostile at the gate—he looks overwhelmed and apologetic.",
            "Key overload features are easy to see on a quick read: executive strain (starting and finishing tasks, holding sequences in mind), administrative avoidance (forms, letters, phone trees), and shutdown when demand spikes. He says he wants to sort things out. Faced with paperwork or ringing services back, he often freezes. Post stays unopened. Calls go unanswered. He dreads being seen as negligent or lazy. Sometimes anxiety spikes and he cannot hold the steps of a phone call in mind.",
            "His flat is harder to keep ordered. Sleep is poor. There are stretches where he shuts down and cannot initiate the smallest task—even meals or washing uniforms. He carries intense guilt, telling workers he is “failing” his children. Support outside immediate services is thin; extended family live at a distance. Previous contacts have raised possible ADHD; he has never been assessed. He links lifelong attention, time-management, and organisation difficulties to unemployment after redundancy, money pressure, and the load of solo parenting. He minimises distress in front of professionals, so he can look more capable than he feels between visits.",
            "On direct questioning he denies wanting to end his life. He speaks of feeling cornered and wishing he could escape. His children anchor him. The picture combines executive-function strain, shame, material stress, and safeguarding questions about the children’s welfare. He is relieved to talk and anxious about new expectations without a viable first step. Cultural and structural factors matter too: how services communicate, and how blame lands on parents under financial strain.",
        ],
        "trainee": "In this exercise, please do not use teaching time asking about missing background details; instead, agree as a group to “fill in the gaps” so you can focus on completing the clinical task.",
        "tasks": [
            "Summarise the main neurodiversity-related and contextual factors that explain Ousmane’s presentation and functioning.",
            "Propose support for both emotional and practical challenges, including any steps that would need liaison outside the counselling hour.",
            "Describe how you would use bifocal vision here: holding realistic limits and environmental strain while also attending to meaning, dignity, and growth within psychosynthesis-informed practice.",
        ],
        "ace": "The case embeds service culture, power, parenting, possible ND, executive overload, and safeguarding so that evidence-informed formulation and ethical boundaries must surface; tasks require integration of psychosynthesis concepts with concrete systemic action, lifting scores on context & ethics, evidence integration, analytical clarity, and originality.",
    },
    {
        "filename": "V2-Emily-Dissociation.docx",
        "title": "ACE-Case Emily — Dissociation",
        "paras": [
            "Emily is twenty-two and in the middle of a degree. She describes anxiety, low mood, and self-harm that has become more frequent over recent months. Deadlines and seminar participation feel impossible on bad weeks; she sits in the library unable to start, then berates herself for wasting fees. Concentration fragments under worry about grades, money, and what happens after graduation. Sleep is patchy. She often eats irregularly, which she dismisses as “not important compared to how I feel in my head.”",
            "She reports repeated episodes of feeling cut off from her body or surroundings, for minutes or longer. Sounds may seem distant; her hands can feel unfamiliar. She loses track of time. She feels frightened when she “comes back,” unsure what happened in the gap. Dissociation can hide risk from her in the moment. It can make self-report uneven from one hour to the next. Risk must be assessed with that in mind.",
            "She missed a counselling appointment without warning. Later she said she had been in a fog for hours and lost track of time and obligation. Lectures are attended less often. She has stopped answering messages from coursemates. She is not certain she can keep herself safe. She is torn about more intensive help, fearing loss of autonomy, leave from study, or being seen as manipulative. Her family value stoicism and “pushing through.” They know only fragments; she expects disappointment or anger if she is fully honest.",
            "Dissociation is central, not peripheral. Stabilisation-first work might include, for example, brief grounding in session, a simple sleep or meal routine to aim for, or one small academic boundary negotiated with clear support (such as a single agreed extension)—illustrations only, not a protocol. The teaching focus is risk in this context, whether current counselling frequency is enough, and what to defer until stability improves. There is no single traumatic index event in the material provided; do not invent trauma history. Work only with what she offers.",
        ],
        "trainee": "In this exercise, please do not use teaching time asking about missing background details; instead, agree as a group to “fill in the gaps” so you can focus on completing the clinical task.",
        "tasks": [
            "Identify dissociative indicators in the narrative and spell out how each might raise or obscure risk.",
            "Decide whether you would continue as planned, intensify support within current services, or escalate (and what escalation means in your setting).",
            "Outline a stabilisation-first plan for the next few sessions, including pacing, focus of interventions, and what you would explicitly not do yet.",
        ],
        "ace": "Subtle dissociation, ambiguous safety, and family context create a clear data-to-decision chain; stabilisation examples anchor the phased task without becoming a manual; tasks target risk, escalation judgement, and phased care—supporting analytical clarity, clinical completeness, evidence integration, and educational value.",
    },
    {
        "filename": "V2-Michael-Risk.docx",
        "title": "ACE-Case Michael — Risk",
        "paras": [
            "Michael is twenty-nine and began counselling after a close relative died about six months ago. Early sessions included some description of the loss. Lately he cancels at short notice. He arrives on time but sits in near silence, or says he has “nothing to report.” When the therapist reflects what is missing, he shrugs or repeats that he is coping. Relatives say he rarely leaves the house. He has stopped seeing friends. He no longer runs or plays music, routines that used to structure his week. They have urged him to see a GP; he refuses, saying talking should be enough or that doctors “pathologise grief.” Day-to-day functioning has thinned: he often skips breakfast or eats one meal late; nights vary from waking around 3 a.m. to sleeping past midday; on harder weeks he may go several days without showering.",
            "There is no clear statement of suicidal intent or plan. Asked directly, he denies thoughts of harming himself; answers stay brief and flat. In session he often says he is “fine” while offering almost no inner detail. His practitioner feels unease that does not map neatly onto a checklist—not enough for emergency services, too much to ignore. Background hints suggest long-standing social withdrawal and possible neurodevelopmental difference; bereavement may sit on top of older isolation and difficulty with change.",
            "Risk judgement here stays genuinely uncertain. Two careful clinicians could read the same facts and disagree on whether to act now or hold the frame. Waiting risks missing deterioration; acting early risks rupture or over-reaction without clearer data. Trainees should spell out what would justify breaching routine confidentiality or seeking urgent input, what can be monitored while the alliance strengthens, and how the therapist’s anxiety might bias the record or the next step. Notes should show what was asked, what was answered, and what remains unknown—not reassurance in place of evidence. The quality of reasoning matters more than one “correct” escalation decision.",
        ],
        "trainee": "In this exercise, please do not use teaching time asking about missing background details; instead, agree as a group to “fill in the gaps” so you can focus on completing the clinical task.",
        "tasks": [
            "Decide whether this situation requires action now (and what action), or a period of continued work within current boundaries, and justify your choice in plain clinical language.",
            "If you choose to wait, define specific observable triggers that would force immediate escalation or re-referral.",
            "Name one supervision question you would bring from this scenario, linking it to either risk reasoning or countertransference.",
        ],
        "ace": "Ambiguous risk, concrete functional detail, and practitioner unease mirror real threshold decisions; tasks force explicit criteria, documentation logic, and supervision use—mapping onto analytical clarity, clinical completeness, context & ethics, and structure without artificial resolution.",
    },
]


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Calibri"
    return p


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_label_para(doc, label, body_text):
    p = doc.add_paragraph()
    r0 = p.add_run(label)
    r0.font.name = "Calibri"
    r0.font.size = Pt(11)
    r0.bold = True
    r1 = p.add_run(body_text)
    r1.font.name = "Calibri"
    r1.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15


def build_one(case):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_heading(doc, case["title"], level=1)

    for para in case["paras"]:
        add_body(doc, para)

    doc.add_paragraph()
    add_label_para(
        doc,
        "Trainee instruction: ",
        case["trainee"],
    )

    doc.add_paragraph()
    p_tasks = doc.add_paragraph()
    r = p_tasks.add_run("Tasks")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    for i, t in enumerate(case["tasks"], 1):
        add_body(doc, f"{i}. {t}")

    doc.add_paragraph()
    add_label_para(
        doc,
        "ACE-Cases score intention: ",
        case["ace"],
    )

    path = OUT_DIR / case["filename"]
    doc.save(str(path))
    print("Wrote", path)


def main():
    for case in CASES:
        build_one(case)


if __name__ == "__main__":
    main()
