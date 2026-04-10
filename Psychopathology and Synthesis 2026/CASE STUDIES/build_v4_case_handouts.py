# -*- coding: utf-8 -*-
"""Build V4 case-study handouts (Word) from V3 clinical content — same folder."""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = Path(__file__).resolve().parent
FONT_PT = 12
INSTRUCTION = (
    "If key information appears to be missing, make reasonable clinical assumptions and proceed."
)

CASES = [
    {
        "filename": "V4-Leah-Shame.docx",
        "title": "Case Leah — Shame",
        "paras": [
            (
                "Leah is twenty-three and self-referred to a low-cost voluntary-sector service after her "
                "GP suggested talking therapy for low mood. The GP letter is thin: brief PHQ scores, no "
                "mention of risk. Her parents, with whom she lives, tell the service she is bright but "
                "“won’t stick at anything.” In their narrative, a young adult should be working full-time "
                "and moving toward independence; her repeated job losses are framed as laziness and as "
                "letting the family down. Leah hears this as proof she is a burden. She has not told her "
                "parents about the medicine-cupboard moment, and she asks whether the service will write "
                "to the GP if she says more—she is unsure what confidentiality means in practice."
            ),
            (
                "Over roughly two years she has left several retail and hospitality jobs within weeks. "
                "Busy environments, shift changes, and managerial check-ins overwhelm her; she goes "
                "absent, then resigns in shame. She spends long hours in her room, sleep reversed, "
                "contact with friends dropped. She describes low energy, self-critical loops, and "
                "post-social rumination. A friend has raised autism or ADHD; Leah relates to sensory "
                "overload and burnout but refuses assessment, saying a label would “make it real” and end "
                "any hope she could still be “normal.” That tension sits in the room: curiosity versus "
                "shame, and fear that naming difference will collapse what little pride she retains."
            ),
            (
                "Risk is not linear. In session two she says she has never self-harmed. In session three, "
                "after a fight at home, she admits standing at the cupboard where paracetamol and "
                "ibuprofen are kept, wanting “everything to stop,” then walking away. She cannot say "
                "how long she stood there or whether she handled the packets. She minimises—“I wouldn’t "
                "really do it”—but her affect goes flat when pressed. She has not told the GP. She says "
                "if the therapist “makes a fuss” her parents will find out and the shame will be "
                "unbearable. Her mother has threatened to stop paying for anything “non-essential” if "
                "she does not find stable work by spring; Leah reads that as conditional love. Part of her "
                "wants structure and honesty; part expects to drop contact with therapy as she has with "
                "jobs. The service has no crisis cover out of hours; you are deciding in a fifty-minute "
                "hour whether this crosses the threshold for same-day GP contact or whether that would "
                "feel like betrayal and end contact. There is no neat formulation yet: passive ideation, "
                "possible impulsive gesture, family pressure, possible neurodevelopmental overlay, "
                "shifting self-report, and incomplete history."
            ),
        ],
        "tasks": [
            (
                "Formulation under uncertainty",
                "Draft a working formulation that holds mood, shame, family context, possible "
                "neurodevelopmental processing differences, and risk—stating explicitly what you do not "
                "know and what would change your view.",
            ),
            (
                "Immediate clinical decision",
                "What do you do after the cupboard disclosure (same day / before the next session)—"
                "including whether to involve the GP, breach confidentiality to parents, or use another "
                "route—and why?",
            ),
            (
                "Longer-term plan and/or supervision reflection",
                "Outline how you would pace the next six to eight sessions if she remains engaged, and "
                "name one question you would take to supervision about either risk or family/system "
                "pressure.",
            ),
        ],
    },
    {
        "filename": "V4-Ousmane-Overload.docx",
        "title": "Case Ousmane — Overload",
        "paras": [
            (
                "Ousmane is thirty-six, a single father of two children aged six and nine. A local family "
                "support worker referred him after missed housing appointments and rising rent arrears. "
                "The school has emailed twice: irregular attendance, incomplete homework, children "
                "sometimes without lunch money or signed trip forms. The head of year has said informally "
                "that if things do not improve they will need to “have a conversation with children’s "
                "services”—not a referral yet, but the edge is visible. Ousmane presents at the school "
                "gate as polite, even charming; staff say he “seems to understand” when they explain what "
                "is needed. The same week he fails to return a call from the housing officer and leaves a "
                "benefits form half-done on the kitchen table for three weeks."
            ),
            (
                "He describes his mind going blank when faced with multi-step tasks, letters he cannot "
                "face opening, and phone menus he cannot hold in sequence once anxiety rises. Sleep is "
                "broken; there are days he cannot initiate meals or washing uniforms though he loves his "
                "children and fears being judged. Previous services noted possible ADHD; he was never "
                "assessed—waiting lists were long and he dropped out of one telephone triage after a "
                "missed callback. He denies suicidal intent but says he feels trapped and sometimes "
                "imagines disappearing; the children are why he stays. He minimises distress in "
                "appointments, then sounds desperate between sessions in a text you are not sure your "
                "service permits you to use clinically."
            ),
            (
                "The ethical tension is sharp: autonomy and trust versus duty of care toward the children "
                "and toward him. If you push liaison with school or social care without consent, you may "
                "lose him; if you stay silent, you may collude with a drift toward formal safeguarding. "
                "NHS ADHD pathways in your area may be twelve to eighteen months; private assessment is "
                "not affordable. His last GP appointment was a rushed ten-minute slot that ended with "
                "“try mindfulness.” He tells you the children ate takeaways four nights one week; the "
                "following week he insists they “mostly had proper meals” and cannot align the details. "
                "You must decide, without a full picture, whether to encourage or facilitate a "
                "safeguarding conversation, a GP re-referral, or a focused stabilisation-only contract "
                "first—and what you document when his account of how bad it is at home shifts between "
                "contacts. Whatever you choose, you will not have everyone’s agreement or complete "
                "information, and the children’s day-to-day safety sits in the same frame as his shame "
                "and exhaustion."
            ),
        ],
        "tasks": [
            (
                "Formulation under uncertainty",
                "Integrate neurodiversity-related and executive-load factors with poverty, parenting load, "
                "shame, and service thresholds—stating what remains unknown.",
            ),
            (
                "Immediate clinical decision",
                "What is your next move after this session (for example consent for school liaison, GP "
                "letter, safeguarding discussion, safety plan only)—with brief rationale and what would "
                "make you revise it?",
            ),
            (
                "Longer-term plan and/or supervision reflection",
                "Sketch a six-session arc that balances practical stabilisation with emotional work, and "
                "one supervision question on either safeguarding boundaries or bifocal vision under "
                "extreme strain.",
            ),
        ],
    },
    {
        "filename": "V4-Emily-Dissociation.docx",
        "title": "Case Emily — Dissociation",
        "paras": [
            (
                "Emily is twenty-two, a second-year undergraduate. Student counselling saw her after her "
                "personal tutor expressed concern about missed seminars and a withdrawn presentation in a "
                "tutorial. The tutor does not know about self-harm; Emily asked that counselling not "
                "contact the department unless she agrees. The university’s student support policy puts "
                "pressure both ways: duty toward student welfare and expectations that courses are informed "
                "when risk is “serious.” Where that line falls is unclear to you."
            ),
            (
                "Emily describes anxiety, low mood, and self-harm that has increased over recent months—"
                "except in the first session she said the last episode was “weeks ago,” and in the third "
                "she could not remember whether she cut three days ago or only thought about it. She "
                "shows you healed marks on her forearm; when you ask how they happened, she stares, then "
                "says she was “probably dissociating” and is not sure of the sequence. Sometimes she "
                "hears her phone buzz from another room but cannot connect the sound to herself. Episodes "
                "of detachment last from minutes to hours; after one she missed a counselling appointment "
                "entirely and later said she lost the afternoon."
            ),
            (
                "She is not certain she can keep herself safe. She fears psychiatric assessment, compulsory "
                "leave, or being labelled attention-seeking. Her family values stoicism; she has told them "
                "little. This week she emailed saying she is “okay now” and asking to cancel the next "
                "session; the tone reads flat. You have no GP collateral. NHS crisis services in your area "
                "typically require explicit intent or a recent attempt for same-day input; her "
                "presentation is fluctuating and partially disclosed."
            ),
            (
                "You must act without a coherent timeline: her memory gaps may reflect dissociation, "
                "minimisation, or shame—or all three. Escalating (breach of confidentiality to the "
                "university, GP, or crisis team) could protect her or could destroy the alliance and push "
                "help underground. Not escalating could leave risk invisible. Her personal tutor has "
                "emailed counselling asking only whether Emily is “engaging”; you have not replied, and "
                "that silence is itself a clinical–ethical fact in a pressured institution. There "
                "is no single traumatic index event in the material; do not invent one. Work only with "
                "what she has said while holding that obscured risk is still risk. The decision you make "
                "today—call her back, write to the GP, inform the tutor, or hold the frame—will be "
                "provisional and contestable, and you will not know for weeks whether it was “right.” "
                "That is the point of the exercise."
            ),
        ],
        "tasks": [
            (
                "Formulation under uncertainty",
                "Link mood, dissociation, academic pressure, family context, and self-harm—explicitly "
                "noting contradictions in her account and what you still cannot know.",
            ),
            (
                "Immediate clinical decision",
                "Given the vague email and unclear recent self-harm, what do you do before the next "
                "scheduled contact (including whether to attempt outreach, involve others, or document "
                "and wait)—and what is the ethical trade-off?",
            ),
            (
                "Longer-term plan and/or supervision reflection",
                "Outline a stabilisation-first sequence for the next few sessions if she remains in "
                "contact, and one supervision question on escalation thresholds or institutional pressure.",
            ),
        ],
    },
    {
        "filename": "V4-Michael-Risk.docx",
        "title": "Case Michael — Risk",
        "paras": [
            (
                "Michael is twenty-nine. He began counselling privately after a close relative died six "
                "months ago. Early sessions included some grief narrative; recently he cancels often, "
                "answers in monosyllables, and says he is “fine.” Relatives report he rarely goes out, has "
                "dropped running and music, and refuses GP contact—he says grief is not an illness. In "
                "session he denies suicidal ideation when asked directly; his voice is flat. Last week he "
                "mentioned in passing that he had looked at evening classes in photography “for when I feel "
                "better,” then said he had not followed up and changed the subject. The future-oriented "
                "snippet sits oddly against withdrawal; it is not the same as clear risk disclosure, but "
                "it contradicts a story of total shutdown without adding clarity. He missed his last "
                "scheduled session, later explaining he had 'lost track of the day,' which fits with his "
                "account but adds to the pattern of inconsistency."
            ),
            (
                "His practitioner’s unease sits below threshold for emergency services: no plan, no intent "
                "stated, no recent attempt disclosed. Yet sleep is fragmented—waking at 3 a.m. or sleeping "
                "past midday—eating is patchy, and hygiene slips on bad weeks. He has lost weight without "
                "intending to; he says he forgets to eat until evening. The contradiction between “I’m "
                "coping” and observable narrowing of life is unresolved. Longstanding social awkwardness "
                "and possible neurodevelopmental difference are background noise, not explained. He works "
                "from home on a contract that barely covers rent; isolation during the day is almost "
                "total, and there is no occupational health route. NHS talking therapies in his area "
                "operate stepped-care thresholds; as a private client he may fall outside those "
                "pathways, which changes what “escalation” means but not whether concern exists. His "
                "sister has offered to attend a session; he refuses, then says “maybe later,” and does "
                "not follow up."
            ),
            (
                "The ethical tension is when concern becomes action without evidence that would satisfy a "
                "crisis team. Autonomy favours continuing the relationship and exploring ambivalence; "
                "duty of care asks whether colluding with minimisation is safe. If you insist on a GP "
                "letter now, he may end therapy; if you do not, you carry the uncertainty alone until "
                "supervision. Documentation must separate facts from inference: what was asked, what he "
                "answered, what remains unknown. Two competent clinicians could disagree on whether to "
                "insist on GP involvement now or to contract for fuller risk review next session—that "
                "disagreement is part of the teaching point, not a flaw in the case. You still have to "
                "choose what you will do before you leave the room, even though the file will stay "
                "incomplete."
            ),
        ],
        "tasks": [
            (
                "Formulation under uncertainty",
                "Integrate bereavement, withdrawal, functional change, and the contradictory future "
                "hint—stating what supports lower immediate risk and what keeps escalation on the table.",
            ),
            (
                "Immediate clinical decision",
                "Do you act now (for example insist on GP contact or shift boundaries) or hold and "
                "contract—and what would flip your choice overnight?",
            ),
            (
                "Longer-term plan and/or supervision reflection",
                "Outline how you would use the next four sessions if he continues, and one supervision "
                "question that separates your anxiety from his data.",
            ),
        ],
    },
]


def add_para(doc, text, bold=False, size_pt=FONT_PT, space_after=Pt(10)):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = space_after
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size_pt)
    r.font.name = "Calibri"


def add_heading_para(doc, text):
    add_para(doc, text, bold=True, size_pt=FONT_PT, space_after=Pt(12))


def build_one(spec):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(FONT_PT)
    style.font.name = "Calibri"

    add_heading_para(doc, spec["title"])
    doc.add_paragraph()

    for block in spec["paras"]:
        add_para(doc, block, space_after=Pt(12))

    add_heading_para(doc, "Trainee instruction")
    add_para(doc, INSTRUCTION, space_after=Pt(14))

    add_heading_para(doc, "Tasks")
    for i, (task_title, task_body) in enumerate(spec["tasks"], start=1):
        add_para(doc, f"{i}. {task_title}", bold=True, space_after=Pt(4))
        add_para(doc, task_body, space_after=Pt(12))

    out = OUT_DIR / spec["filename"]
    doc.save(str(out))
    return out


def main():
    for spec in CASES:
        path = build_one(spec)
        print("Wrote", path)


if __name__ == "__main__":
    main()
