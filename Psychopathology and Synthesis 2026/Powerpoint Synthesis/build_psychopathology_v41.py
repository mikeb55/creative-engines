# -*- coding: utf-8 -*-
"""Build Psychopathology & Synthesis deck: large type, Calibri, rewritten speaker notes.

Versioning (Psychopathology & Synthesis line):
  This script exports V4.1. The next increment should be V4.2 (not V5.x or V7.x).

Title slide:
  Load PET PowerPoint Title page.pptx, keep only slide 1, set title/subtitle text only
  (preserve PET master branding—do not force Calibri on that slide).
"""
from pptx import Presentation
from pptx.util import Pt

# Official PET title layout (slide 1); includes same content layouts as teaching slides.
TITLE_SOURCE = r"C:\Users\mike\Documents\PET\Psychopathology and Synthesis 2026\Powerpoint Synthesis\PET PowerPoint Title page.pptx"
OUT = r"C:\Users\mike\Documents\PET\Psychopathology and Synthesis 2026\Powerpoint Synthesis\V4.1-Psychopathology-Synthesis.pptx"

# Typography (room visibility)
TITLE_CONTENT_PT = 40
BODY_CONTENT_PT = 28
SECTION_TITLE_PT = 44
SECTION_SUB_PT = 26
TITLE_ONLY_PT = 44
REFERENCES_TITLE_PT = 36
REFERENCES_BODY_PT = 17


def trim_to_first_slide(prs):
    """Remove sample slides; keep only slide 1 (PET branded title)."""
    while len(prs.slides) > 1:
        rId = prs.slides._sldIdLst[1].rId
        prs.part.drop_rel(rId)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[1])


def apply_pet_branded_title(slide, title, subtitle, note_points):
    """Set placeholder text only; do not override theme fonts on the PET title slide."""
    for shape in slide.shapes:
        if not shape.is_placeholder:
            continue
        idx = shape.placeholder_format.idx
        if idx == 0:
            shape.text = title
        elif idx == 1:
            shape.text = subtitle
    add_notes_bullets(slide, note_points)


def add_notes_bullets(slide, points):
    """2–3 short bullet lines for speaker notes (scannable)."""
    points = [p.strip() for p in points if p.strip()][:3]
    lines = [f"• {p}" for p in points]
    notes = slide.notes_slide
    tf = notes.notes_text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        if i == 0:
            para = tf.paragraphs[0]
            para.text = line
        else:
            para = tf.add_paragraph()
            para.text = line
        para.font.name = "Calibri"
        para.font.size = Pt(18)


def style_title_shape(shape, size_pt, bold=True):
    if shape is None or not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        p.font.name = "Calibri"
        p.font.size = Pt(size_pt)
        p.font.bold = bold
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(size_pt)
            r.font.bold = bold


def style_body_shape(shape, size_pt, bold=False):
    if shape is None or not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        p.font.name = "Calibri"
        p.font.size = Pt(size_pt)
        p.font.bold = bold
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(size_pt)
            r.font.bold = bold


def bullets_from_lines(body_shape, lines, body_pt):
    tf = body_shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.name = "Calibri"
        p.font.size = Pt(body_pt)
        p.font.bold = False
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(body_pt)


def slide_title_content(prs, title, lines, note_points, title_pt=TITLE_CONTENT_PT, body_pt=BODY_CONTENT_PT):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    style_title_shape(slide.shapes.title, title_pt, bold=True)
    ph = slide.placeholders[1]
    bullets_from_lines(ph, lines, body_pt)
    style_body_shape(ph, body_pt)
    add_notes_bullets(slide, note_points)
    return slide


def slide_section(prs, title, subtitle_lines, note_points, sub_pt=SECTION_SUB_PT):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text = title
    style_title_shape(slide.shapes.title, SECTION_TITLE_PT, bold=True)
    if subtitle_lines:
        ph = slide.placeholders[1]
        bullets_from_lines(ph, subtitle_lines, sub_pt)
        style_body_shape(ph, sub_pt)
    add_notes_bullets(slide, note_points)
    return slide


def slide_title_only(prs, title, note_points, title_pt=TITLE_ONLY_PT):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    style_title_shape(slide.shapes.title, title_pt, bold=True)
    add_notes_bullets(slide, note_points)
    return slide


def main():
    prs = Presentation(TITLE_SOURCE)
    trim_to_first_slide(prs)

    apply_pet_branded_title(
        prs.slides[0],
        "Psychopathology & synthesis",
        "Psychosynthesis & Education Trust\nPostgraduate weekend · 11–12 April 2026",
        [
            "State the aim in one sentence: using psychopathology and psychosynthesis together in accountable clinical practice.",
            "Point to the next slide for typical PET timings; adjust to your contract and room.",
            "Name one practical thread for the weekend: formulation, risk assessment, and what trainees might do after Monday.",
        ],
    )

    slide_title_content(
        prs,
        "Weekend rhythm & time",
        [
            "Saturday & Sunday: 10:00–17:00 (UK)",
            "Each day: about 15 min meditation; optional 15 min check-in; teaching; 20 min morning break; lunch (often 13:00–14:00); 20 min afternoon break.",
            "Facilitators: allocate minutes so meditation, check-in, teaching, breaks, and lunch still fit the day.",
        ],
        [
            "Clarify that the deck’s labels follow an earlier PET version: “Day one — afternoon” comes after the mid-morning break and before the lunch slide, not necessarily after 14:00.",
            "List the major runs for your own timetable: framing through Leah; neurodiversity through Ousmane; day two trauma and dissociation; then risk, ambiguity, and closure.",
            "Remind trainees that breaks are scheduled teaching time for regulation and attention, not optional padding.",
        ],
    )

    slide_title_content(
        prs,
        "How we learn this weekend",
        [
            "You bring substantial professional and life experience; the level is postgraduate formation, not introductory survey.",
            "We prioritise clinical judgement, realistic uncertainty, and decisions you could defend in supervision.",
            "Pairs and short reflection are methods for application, the same skills you use preparing for placement.",
        ],
        [
            "Give one example of transferable skill: risk assessment from nursing, teaching, or safeguarding work maps onto counselling thresholds with different paperwork.",
            "Clarify that “clinical judgement” means reasons you can name (what you saw, asked, recorded), not preference without evidence.",
            "Optional invitation: connect theory to a past role only if the trainee chooses; no requirement to disclose personal mental health history.",
        ],
    )

    slide_title_content(
        prs,
        "Day one — opening",
        [
            "Meditation: about 15 minutes (PET tradition).",
            "Optional check-in: about 15 minutes — “How do you feel about this material?”",
            "Then begin the day one — morning teaching block.",
        ],
        [
            "Say who leads meditation and what you expect on cameras, doors, or movement.",
            "Offer an alternative to a full group check-in: pair share for two minutes, or a single written sentence, if the group is large.",
            "Name when the morning break will be so trainees can pace attention.",
        ],
    )

    slide_section(
        prs,
        "Day one — morning",
        ["Framing · psychopathology · risk · case (Leah) · psychosynthesis threads"],
        [
            "Preview the four anchors: how we use diagnosis, risk with Leah, then psychosynthesis language for parts and integration.",
            "Give an approximate end time for this block before you start, if you know it.",
            "Position the case as professional task practice, not entertainment or gossip.",
        ],
    )

    slide_title_content(
        prs,
        "Framing the work",
        [
            "Psychopathology offers shared clinical language.",
            "It supports formulation—not identity collapse.",
            "Meaning and risk must be held together.",
            "Therapists tolerate uncertainty without freezing or rushing.",
        ],
        [
            "Clarify that diagnosis is descriptive shorthand for patterns, not the whole person.",
            "Give one example of when shared language with a GP or psychiatrist matters (e.g. agreeing what “low mood” and “risk” mean in a shared note).",
            "Remind trainees that uncertainty does not remove responsibility for action when a risk threshold is crossed.",
        ],
    )

    slide_title_content(
        prs,
        "Working with psychopathology",
        [
            "Diagnostic labels describe patterns—not people.",
            "Formulation integrates symptoms, history, and context.",
            "Clients arrive shaped by prior labels and systems.",
            "Think clinically without collapsing into diagnosis as destiny.",
        ],
        [
            "Name typical sources of prior labels: CAMHS, CMHT, occupational health, or private assessment, and how those letters shape client expectation.",
            "Contrast a formulation you revise in supervision with a label used as a final explanation of character.",
            "Ask for one example each: a time a diagnosis helped access support, and a time it narrowed what services listened for.",
        ],
    )

    slide_title_content(
        prs,
        "Risk — core principles",
        [
            "Risk is dynamic and can escalate quickly.",
            "Intensity does not equal level of danger.",
            "Quiet clients may carry significant risk.",
            "Direct, respectful questions reduce risk when asked with care.",
        ],
        [
            "List factors that commonly shift risk: isolation, substance use, sleep loss, humiliation, or recent loss, not only loud distress.",
            "Give a concrete question sequence: thoughts, intent, plan, means, and protective factors, without dramatising the room.",
            "Remind trainees that your training or agency risk protocol applies here; cite it by name if you have one.",
        ],
    )

    slide_title_only(
        prs,
        "Case study · Shame — “Leah”",
        [
            "Present the written case or handout in full; allow a short silence before discussion starts.",
            "Keep discussion on professional response and formulation, not on diagnosing anyone present.",
            "Link shame to missed disclosure: clients may minimise risk when they expect judgment.",
        ],
    )

    slide_title_content(
        prs,
        "Task · Leah (1)",
        [
            "Clarify thoughts, intent, and access to means.",
            "Develop formulation: shame, avoidance, context.",
            "Prioritise the first two sessions and pacing.",
        ],
        [
            "Expand “access to means” with a mundane example: stockpiled medication, a specific location, or a planned time.",
            "Formulation here means a short narrative linking affect, relationships, and maintaining factors, not a longer essay.",
            "First two sessions: contract, safety baseline, and what can wait until alliance exists.",
        ],
    )

    slide_title_content(
        prs,
        "Task · Leah (2)",
        [
            "Match structure to client capacity.",
            "Reflective prompt: What would you not do here, and why?",
        ],
        [
            "Give examples of mismatch: long interpretive monologues when the client is flooded, or rigid homework when shame blocks follow-through.",
            "The “not do” prompt surfaces common errors: rescuing, premature interpretation, or closing risk too early to soothe anxiety.",
            "Bridge to the next slide: meaning can be discussed without lowering safeguarding standards.",
        ],
    )

    slide_title_content(
        prs,
        "Psychosynthesis bridge · Meaning alongside risk",
        [
            "Suffering can be meaningful without minimising danger.",
            "Non-identification supports seeing patterns without fusion or contempt.",
            "Risk decisions stay ethical and ordinary—not heroic.",
        ],
        [
            "Distinguish meaning-making after assessment from using meaning to avoid asking about suicide.",
            "Firman and Gila’s work on identification helps explain why therapists over-identify with rescuer or expert roles under pressure.",
            "Nelson’s integration theme supports careful work with “split” experiences without romanticising breakdown as growth.",
        ],
    )

    slide_title_content(
        prs,
        "Psychosynthesis perspective",
        [
            "Distress can reflect fragmentation of parts (subpersonalities).",
            "An observing centre supports disidentification where appropriate.",
            "Containment is required before exploration.",
            "Integration develops gradually—over time and relationship.",
        ],
        [
            "Map subpersonalities to everyday therapy language: inner critic, frightened child, compliant self.",
            "Disidentification in session often starts with naming affect and simple grounding, not immediate transpersonal claims.",
            "Sequence matters: stabilisation and relationship before large interpretations or deep regression work.",
        ],
    )

    slide_title_only(
        prs,
        "Morning break · 20 minutes",
        [
            "State the return time and where refreshments are.",
            "Next block covers neurodiversity through Ousmane and practical strategies.",
            "If someone needs pastoral follow-up, name the PET or course contact rather than extending the break into informal therapy.",
        ],
    )

    slide_section(
        prs,
        "Day one — afternoon",
        ["Neurodiversity · clinical challenges · case (Ousmane) · practical work"],
        [
            "Preview misreadings of neurodivergence as motivation problems and the role of environment in formulation.",
            "Estimate time to lunch if your Saturday schedule differs from the slide order.",
            "Point forward to safeguarding and children in Ousmane’s task where relevant.",
        ],
    )

    slide_title_content(
        prs,
        "Neurodiversity — clinical view",
        [
            "Difference in processing rather than deficit-only framing.",
            "Executive function affects organisation and action.",
            "Sensory overload drives overwhelm and shutdown.",
            "Masking hides difficulty and increases exhaustion.",
        ],
        [
            "Executive function includes working memory, inhibition, and planning; difficulty reflects load, not a moral label.",
            "Sensory examples: fluorescent lights, overlapping speech in reception, unpredictable schedule changes.",
            "Masking explains late diagnosis in some groups: the client looked fine in brief contacts.",
        ],
    )

    slide_title_content(
        prs,
        "Clinical challenges",
        [
            "Difference is often misread as lack of motivation.",
            "Shame develops through repeated failure narratives.",
            "Function varies depending on environment.",
            "Work requires flexibility without false certainty.",
        ],
        [
            "Workplace example: performance management language may blame the person before adjusting tasks or environment.",
            "Shame often follows years of being called careless; acknowledge effort before cognitive restructuring.",
            "Formulation updates when the environment changes: new job, new relationship, or new housing.",
        ],
    )

    slide_title_content(
        prs,
        "Psychosynthesis bridge · Adaptation, not pretend normality",
        [
            "Adjustment of environment and pace can be deeply ethical—not “lowering standards.”",
            "Integration includes honest limits; wellness does not require looking neurotypical.",
            "Avoid idealising neurodivergence or collapsing it into tragedy.",
        ],
        [
            "Reasonable adjustments in study or work are legal and clinical tools, not excuses.",
            "Firman and Gila’s attention to context fits adjusting pace, structure, and relationship alongside inner work.",
            "One coherent stance across the week beats praising neurodivergence as gift in one session and treating it as disorder in the next.",
        ],
    )

    slide_title_only(
        prs,
        "Case study · Neurodiversity — “Ousmane”",
        [
            "Introduce the case at the same pace as Leah: facts first, then tasks.",
            "Highlight stressors that are systemic (housing, work hours, childcare) as well as internal.",
            "Note safeguarding for dependents when the slide deck flags children.",
        ],
    )

    slide_title_content(
        prs,
        "Task · Ousmane (1)",
        [
            "Formulate neurodiversity, stress, and environment together.",
            "Identify safeguarding risks for children where relevant.",
            "Prioritise practical next steps.",
        ],
        [
            "Show how one formulation sentence links ND traits, external pressure, and coping.",
            "Safeguarding: what a child sees, hears, or depends on when the parent is overwhelmed.",
            "Practical next steps might include GP, employer, school liaison, or benefits—stay within role boundaries.",
        ],
    )

    slide_title_content(
        prs,
        "Task · Ousmane (2)",
        [
            "Focus on what might change this week.",
            "Pair discussion: one environmental change; one relational stance you would take.",
        ],
        [
            "Environmental change examples: a quieter appointment time, written instructions, or fewer simultaneous demands.",
            "Relational stance examples: slower pace, explicit structure for the hour, or validating fatigue before problem-solving.",
            "Documentation: one line you would want in the record if another professional read it tomorrow.",
        ],
    )

    slide_title_content(
        prs,
        "Working practically",
        [
            "Use structure, routine, and predictability where helpful.",
            "Reduce overwhelm before increasing insight.",
            "Adjust environment where possible.",
            "Support agency and functional change—without rushing “depth.”",
        ],
        [
            "Ordinary tools count: calendars, reminders, graded tasks, and clear session agendas.",
            "Insight-first sequences often fail when arousal is high; sequence stabilisation before interpretation.",
            "Depth-oriented psychosynthesis work usually follows practical stability, not instead of it.",
        ],
    )

    slide_title_only(
        prs,
        "Lunch · 1 hour",
        [
            "Lunch is often scheduled 13:00–14:00; confirm with your cohort.",
            "This deck places day-one afternoon content before this slide; day two follows after lunch in the file order.",
            "If your Saturday continues after lunch, resume at the slide that matches your live plan.",
        ],
    )

    slide_title_content(
        prs,
        "Day two — opening",
        [
            "Meditation: about 15 minutes.",
            "Check-in: about 15 minutes — “What did yesterday bring up for you?”",
            "Then day two — morning block: trauma, dissociation, Emily.",
        ],
        [
            "Overnight processing is common after risk and shame material; sharing detail remains optional.",
            "If a trainee is acutely distressed, use PET pastoral or safeguarding routes you already have, not improvised group processing.",
            "Preview that the morning block ends at the afternoon break before risk cases.",
        ],
    )

    slide_section(
        prs,
        "Day two — morning",
        ["Trauma · dissociation · case (Emily) · neuroplasticity · clinical response"],
        [
            "List topics: nervous system effects of trauma, dissociation, Emily case, neuroplasticity, and phased clinical response.",
            "Give an approximate finish time before the afternoon break.",
            "Note that dissociation content links to later risk slides on David and Michael.",
        ],
    )

    slide_title_content(
        prs,
        "Trauma — core understanding",
        [
            "Trauma is held in the nervous system—not only in “the story.”",
            "It alters regulation, perception, and relating.",
            "Triggers can produce rapid state changes.",
            "Clients may shift between arousal and shutdown.",
        ],
        [
            "Contrast narrative memory with body-based reactivity that appears before words organise.",
            "Trigger example: a tone of voice or smell linked to earlier threat.",
            "Shutdown can look like compliance or flat affect rather than obvious panic.",
        ],
    )

    slide_title_content(
        prs,
        "Dissociation",
        [
            "Often a protective response to overwhelming experience.",
            "May involve loss of continuity or presence.",
            "Can be subtle and easily missed in clinic.",
            "Strongly linked to increased risk—assess carefully.",
        ],
        [
            "Subtle signs: time slips, depersonalisation, “going away” mid-sentence, or inconsistent recall of session content.",
            "Distinguish ordinary distraction from dissociative gap when risk is present.",
            "Risk rises when continuity fails and self-harm or suicide plans are not fully accessible to the client’s conscious report.",
        ],
    )

    slide_title_content(
        prs,
        "Psychosynthesis bridge · Containment before exploration",
        [
            "Rapid deepening is not always kind when stability is fragile.",
            "Stabilisation and relationship come first.",
            "Transpersonal material is not a shortcut past fear or trauma.",
        ],
        [
            "Grof on spiritual emergency and Clarke on psychosis and spirituality are background reading on extreme states, not step-by-step methods for beginners.",
            "Nelson’s theme of integration cautions against forcing “split” material into insight before safety exists.",
            "Stay within competence, contract, and supervision for trauma and dissociation work.",
        ],
    )

    slide_title_only(
        prs,
        "Case study · Dissociation — “Emily”",
        [
            "Read the case at a steady pace; dissociation deserves the same dignity as any other presentation.",
            "Invite hypotheses about what an observer might miss in a single session.",
            "Foreshadow that the next slides ask for risk and escalation decisions.",
        ],
    )

    slide_title_content(
        prs,
        "Task · Emily (1)",
        [
            "Assess risk alongside dissociation.",
            "Decide whether to escalate support.",
            "Prioritise stabilisation and safety.",
        ],
        [
            "Escalation might mean crisis team, GP, designated safeguarding lead, or your agency’s on-call route—name what your context actually uses.",
            "Stabilisation includes grounding, clear boundaries, frequency of contact, and sometimes pausing exploratory work.",
            "Risk and dissociation together require written rationale when you choose not to escalate.",
        ],
    )

    slide_title_content(
        prs,
        "Task · Emily (2)",
        [
            "Question reliability of self-report—without contempt.",
            "Reflective prompt: What would change your mind overnight about risk?",
        ],
        [
            "Unreliable report can mean shame, fear, or amnesic gaps, not manipulation; language matters in notes.",
            "The overnight prompt targets concrete indicators: new means, substance use, relationship rupture, or psychotic content.",
            "Link forward: neuroplasticity explains why change is slow and relational.",
        ],
    )

    slide_title_content(
        prs,
        "Psychosynthesis bridge · Self, woundedness, and pace",
        [
            "Depth is compatible with safety: smaller steps, repeated contact.",
            "Meaning may emerge without forcing narrative coherence immediately.",
            "Therapist pacing shapes how much arousal the hour can hold.",
        ],
        [
            "Firman and Gila’s framing allows depth and spirit without promising linear recovery.",
            "Small repeatable interventions (sleep routine, one safe person, one grounding skill) match the next slide on neuroplasticity.",
            "Slower pace is a clinical choice when integration would flood the client.",
        ],
    )

    slide_title_content(
        prs,
        "Neuroplasticity",
        [
            "Repetition strengthens emotional patterns—for better and worse.",
            "Trauma can reinforce rigid survival responses.",
            "Insight alone rarely changes embodied behaviour.",
            "Change often requires repeated relational experience over time.",
        ],
        [
            "Negative example: repeated self-criticism deepens automatic shame loops.",
            "Positive example: predictable empathic responses in therapy gradually shift expectation of others.",
            "This is why homework and between-session practice matter alongside interpretation.",
        ],
    )

    slide_title_content(
        prs,
        "Clinical response",
        [
            "Stabilise before deeper exploration.",
            "Use pacing to prevent dysregulation.",
            "Reduce overwhelm to support engagement.",
            "Monitor risk as patterns shift—especially after stabilisation begins.",
        ],
        [
            "When clients feel safer, new material may surface; risk assessment is ongoing, not one-off.",
            "Pacing includes shorter interventions, more checking in, and fewer interpretations per hour when arousal is high.",
            "Link to day one practical work: external structure supports the same nervous system.",
        ],
    )

    slide_title_only(
        prs,
        "Afternoon break · 20 minutes",
        [
            "Afternoon block focuses on acute risk, ambiguity, and integration.",
            "State return time clearly.",
            "Briefly preview David and Michael as decision cases, not new theory blocks.",
        ],
    )

    slide_section(
        prs,
        "Day two — afternoon",
        ["Risk (David) · ambiguity (Michael) · integration · closing"],
        [
            "This block applies yesterday’s principles to escalation, documentation, and uncertain presentations.",
            "Estimate time to the closing reflection and references.",
            "Remind trainees of local emergency and safeguarding numbers if you display them institutionally.",
        ],
    )

    slide_title_only(
        prs,
        "Case study · Risk — “David”",
        [
            "Present facts with clarity; risk teaching works best without theatrical emphasis.",
            "Separate legal or statutory steps from counselling ethics: name what your jurisdiction expects.",
            "Invite initial gut responses, then move to structured decision criteria.",
        ],
    )

    slide_title_content(
        prs,
        "Task · David (1)",
        [
            "Decide: continue or escalate.",
            "Act immediately on current risk where indicated.",
            "Document clearly and justify decisions.",
        ],
        [
            "Continue versus escalate: contrast more frequent sessions with same worker versus duty psychiatrist, AMHP, or 999.",
            "Immediate action examples: removing means with agreement, contacting agreed third party, or emergency services when threshold met.",
            "Documentation: time, questions asked, answers, rationale, and who was informed.",
        ],
    )

    slide_title_content(
        prs,
        "Task · David (2)",
        [
            "Consider impact of trauma on disclosure.",
            "Pair: one sentence you would write in contemporaneous notes.",
        ],
        [
            "Trauma may delay or fragment disclosure; absence of a coherent story does not equal absence of risk.",
            "Notes should be factual enough for another professional to follow your reasoning tomorrow.",
            "Next slide generalises risk decisions beyond this single case.",
        ],
    )

    slide_title_content(
        prs,
        "Risk decisions",
        [
            "Balance autonomy with duty of care.",
            "Escalate when risk becomes active.",
            "Use supervision for complex decisions.",
            "Clear documentation protects clients and practitioners.",
        ],
        [
            "Autonomy includes collaborating on safety plans when capacity allows; duty of care overrides confidentiality at defined thresholds.",
            "Supervision is for proportionate decisions, not only for mistakes after the fact.",
            "Records protect trainees later if a complaint asks what was known and when.",
        ],
    )

    slide_title_only(
        prs,
        "Case study · Ambiguity — “Michael”",
        [
            "Ambiguous risk is common in mature practice; the task is structured judgement, not instant certainty.",
            "Ask what information is missing and whether it can be obtained safely.",
            "Position Michael after David to contrast clear escalation with grey-zone management.",
        ],
    )

    slide_title_content(
        prs,
        "Task · Michael",
        [
            "Decide whether uncertainty requires action now.",
            "Define what would trigger escalation.",
            "Consider risks of acting versus waiting.",
            "Notice therapist anxiety in decision-making—without shame.",
        ],
        [
            "Action now might mean consultation, not always breach of confidentiality; name your options.",
            "Triggers for escalation should be written: specific behaviours, statements, or third-party reports.",
            "Clinician anxiety reflects your own risk history as well as client data; supervision helps sort the two. In a very large group, work one scenario through to a single decision.",
        ],
    )

    slide_title_content(
        prs,
        "Psychosynthesis bridge · Formulation and clinical work",
        [
            "Psychosynthesis offers parts, Self, meaning, and context—not a bypass around ethics or diagnosis.",
            "Integration without idealisation: clients remain in real systems.",
            "Relationship and environment matter as much as inner process.",
        ],
        [
            "Tie Nelson on integration of split experience to formulation that includes biology, psychology, and social context.",
            "Firman and Gila keep spiritual depth inside psychological ethics rather than above ordinary safeguarding.",
            "Optional: critical medicalisation readings (e.g. Boyle in Madness contested) only if you teach that strand and have full citation details.",
        ],
    )

    slide_title_content(
        prs,
        "Integration",
        [
            "Hold meaning and risk simultaneously.",
            "Recognise limits of the therapeutic role.",
            "Decision-making is core to ethical practice.",
            "Use reflection and supervision consistently.",
        ],
        [
            "Limits of role: you are not police, GP, or family; you know what your contract permits.",
            "Decision-making includes when not to act as well as when to escalate.",
            "One takeaway: identify one supervision question you will bring from this weekend.",
        ],
    )

    slide_title_content(
        prs,
        "Closing reflection",
        [
            "Short pause or silence before you close, if you use one.",
            "Optional: one word from the weekend, or pass; declining is fine.",
            "Thank the group for attention to difficult material.",
        ],
        [
            "If time is short, cut this slide before cutting risk or trauma content.",
            "Avoid requiring positive closure; fatigue and ambivalence are legitimate outcomes.",
            "Direct anyone who needs ongoing support to course or PET procedures.",
        ],
    )

    ref_body = [
        "Nelson, J.E. (1994) Healing the split: Integrating spirit into our understanding of the mentally ill. Albany, NY: State University of New York Press.",
        "Firman, J. and Gila, A. (2002) Psychosynthesis: A psychology of the spirit. Albany, NY: State University of New York Press.",
        "Firman, J. and Gila, A. (1997) The primal wound: A transpersonal view of trauma, addiction, and growth. Albany, NY: State University of New York Press.",
        "Grof, S. and Grof, C. (eds) (1989) Spiritual emergency: When personal transformation becomes a crisis. [Publication details to confirm.]",
        "Clarke, I. (ed.) (2001) Psychosis and spirituality: Exploring the new frontier. [Publication details to confirm.]",
        "Clarke, I. (ed.) (2010) Psychosis and spirituality: Consolidating the new paradigm. [Publication details to confirm.]",
    ]

    slide_title_content(
        prs,
        "Selected references",
        ref_body,
        [
            "The first three entries are complete as given; match your programme’s Harvard conventions for punctuation.",
            "Grof and Clarke lines lack verified publisher and place in the master list; confirm in the library catalogue before student essays.",
            "Boyle in Madness contested is optional critical reading; confirm full details before adding to a formal bibliography.",
        ],
        title_pt=REFERENCES_TITLE_PT,
        body_pt=REFERENCES_BODY_PT,
    )

    prs.save(OUT)
    print("Saved", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
