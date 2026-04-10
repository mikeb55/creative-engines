# -*- coding: utf-8 -*-
"""Build V4.4 Psychopathology & Synthesis — full weekend, V4 case handouts, PET template."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
CASE_DIR = ROOT / "CASE STUDIES"
TITLE_SOURCE = SCRIPT_DIR / "PET PowerPoint Title page.pptx"
OUT = SCRIPT_DIR / "V4.4-Psychopathology-Synthesis.pptx"

TITLE_PT = 32
BODY_PT = 22
BODY_FONT = "Arial Narrow"
TITLE_FONT = "Calibri"
REF_BODY_PT = 16
NOTES_PT = 11

TRAINEE_ASSUME = (
    "If key information appears to be missing, make reasonable clinical assumptions and proceed."
)
INCOMPLETE = "Work with incomplete and contradictory information."


def trim_to_first_slide(prs):
    while len(prs.slides) > 1:
        rId = prs.slides._sldIdLst[1].rId
        prs.part.drop_rel(rId)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[1])


def apply_branded_title(slide, title, subtitle, notes):
    for shape in slide.shapes:
        if not shape.is_placeholder:
            continue
        idx = shape.placeholder_format.idx
        if idx == 0:
            shape.text = title
        elif idx == 1:
            shape.text = subtitle
    add_notes(slide, notes)


def add_notes(slide, points, timing=None):
    """Speaker notes: optional timing line first, then up to five content bullets (max six lines total)."""
    ordered = []
    if timing:
        ordered.append(timing.strip())
    ordered.extend(p.strip() for p in points if p.strip())
    ordered = ordered[:6]
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(ordered):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = f"• {line}"
        para.font.name = "Calibri"
        para.font.size = Pt(NOTES_PT)


def style_content_title_body(slide):
    """Calibri 32 title + Arial Narrow 22 body, centred, vertically centred (slide > 0)."""
    t = slide.shapes.title
    if t.has_text_frame:
        tf = t.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in tf.paragraphs:
            p.font.name = TITLE_FONT
            p.font.size = Pt(TITLE_PT)
            p.font.bold = True
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name = TITLE_FONT
                r.font.size = Pt(TITLE_PT)
                r.font.bold = True
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            continue
        if not ph.has_text_frame:
            continue
        tf = ph.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in tf.paragraphs:
            p.font.name = BODY_FONT
            p.font.size = Pt(BODY_PT)
            p.font.bold = True
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name = BODY_FONT
                r.font.size = Pt(BODY_PT)
                r.font.bold = True


def slide_tc(prs, title, bullets, notes, timing=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
    style_content_title_body(slide)
    add_notes(slide, notes, timing=timing)
    return slide


def slide_to(prs, title, notes, timing=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    tf = slide.shapes.title.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in tf.paragraphs:
        p.font.name = TITLE_FONT
        p.font.size = Pt(TITLE_PT)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.name = TITLE_FONT
            r.font.size = Pt(TITLE_PT)
            r.font.bold = True
    add_notes(slide, notes, timing=timing)
    return slide


def slide_blank(prs, notes, timing=None):
    """PET template ‘Blank’ layout (master branding only) for a light visual pause."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_notes(slide, notes, timing=timing)
    return slide


def slide_sec(prs, title, bullets, notes, timing=None):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text = title
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
    style_content_title_body(slide)
    add_notes(slide, notes, timing=timing)
    return slide


def main():
    prs = Presentation(str(TITLE_SOURCE))
    trim_to_first_slide(prs)

    apply_branded_title(
        prs.slides[0],
        "Psychopathology & synthesis",
        "Psychosynthesis & Education Trust\nPostgraduate weekend · 11–12 April 2026",
        [
            "Timing — Case work (teaching slides + tasks, venue-dependent): Leah ~40–45 min; Ousmane ~50–55 min; Emily ~60 min; Michael ~60–65 min; plus contracted breaks and lunch.",
            "The weekend links diagnostic language with formulation, risk assessment, dissociation, and psychosynthesis concepts used as clinical frameworks rather than as spiritual bypass.",
            "Cases match the V4 Word handouts in the CASE STUDIES folder (V4-Leah-Shame.docx through V4-Michael-Risk.docx); they are anonymised teaching material, not live records.",
            "Risk teaching assumes you map decisions to your own agency, NHS interface, and professional indemnity requirements.",
        ],
    )

    slide_tc(
        prs,
        "Weekend rhythm and time",
        [
            "Saturday and Sunday: 10:00–17:00 (UK), as contracted for this module.",
            "Each day includes about fifteen minutes meditation, an optional fifteen-minute check-in, teaching blocks, a twenty-minute morning break, lunch (often 13:00–14:00), and a twenty-minute afternoon break.",
            "In this file, day-one afternoon teaching appears before the lunch slide, then day two; live timetables may reorder blocks while preserving total teaching and break time.",
        ],
        [
            "Attention and memory for risk material decline across long unbroken stretches; breaks are part of safe design for trauma and suicide content.",
            "Net scheduled teaching time is always less than seven hours per day once openings, lunch, and breaks are subtracted.",
            "The same slide sequence can be split across a live timetable if your venue uses a different lunch placement.",
        ],
        timing="Timing — Map these intervals onto your cohort paperwork; confirm lunch clock locally (often 13:00–14:00).",
    )

    slide_tc(
        prs,
        "Teaching frame: authority, containment, engagement",
        [
            "Authority: clear contracts, accountable risk questions, and decisions you could defend to a regulator or employer.",
            "Containment: pacing, boundaries on depth, and stabilisation before exploratory work where dissociation or shame dominates.",
            "Engagement: active formulation tasks, uncertainty tolerated, and reasoning made visible rather than impressionistic.",
        ],
        [
            "ACE here names the holding structure for postgraduate clinical training, not a separate therapy model.",
            "Containment overlaps with psychosynthesis emphasis on relationship and sequence before interpretation.",
            "Engagement matches ACE-Cases design: each vignette ends in formulation, immediate decision, and longer-term or supervision reflection.",
        ],
        timing="Timing — Early Day 1: keep this framing compact so Leah (40–45 min including tasks) still fits before the morning break.",
    )

    slide_tc(
        prs,
        "Postgraduate learning posture",
        [
            "You bring substantial prior professional and life experience; the standard is formation, not introductory recall.",
            "Clinical judgement means reasons tied to observations, questions asked, and record entries—not preference without evidence.",
            "Pairs and small discussion translate theory into decisions you might make under your contract and supervision arrangement.",
        ],
        [
            "Transferable skills from healthcare, education, or justice settings alter what feels familiar but not what counts as proportionate risk action in counselling.",
            "Documentation habits from other roles need aligning with counselling confidentiality and data protection.",
            "Personal mental health disclosure is never required for participation.",
        ],
        timing="Timing — Still Day 1 AM: if you are past ~45 minutes from day start and still in pre-Leah material, shorten discussion to protect the Leah window.",
    )

    slide_tc(
        prs,
        "Diagnosis and formulation",
        [
            "Diagnosis functions as a working hypothesis: a shorthand for clustering symptoms and functioning, open to revision as context emerges.",
            "Formulation is evolving and contextual: it links history, relationships, culture, environment, and risk in a narrative you update in supervision.",
            "Neither replaces the other; both can sit in the same file if the reasoning for each is explicit.",
        ],
        [
            "NHS and third-sector services often require diagnostic codes for access; that administrative fact does not make the code the whole person.",
            "A strong formulation states what would falsify it—for example new risk information or a change in housing or employment.",
            "Trainees sometimes over-weight diagnosis when GP letters are the only collateral; formulation corrects that bias.",
        ],
        timing="Timing — Day 1 AM: diagnosis versus formulation feeds directly into Leah; allow a few minutes for questions before the case title slide.",
    )

    slide_tc(
        prs,
        "Day one — opening",
        [
            "Meditation of about fifteen minutes follows PET convention for starting teaching days.",
            "Optional check-in of about fifteen minutes: “How do you feel about this material?”",
            "Then move into the day one morning block: framing, risk principles, and the Leah case.",
        ],
        [
            "Brief contemplative openings are common before psychologically dense material and may lower physiological arousal for some participants.",
            "Check-in remains bounded; disclosure of personal mental health history is not a requirement for learning.",
            "The morning block stacks formulation, shame, and passive risk before neurodiversity and overload in the afternoon.",
        ],
        timing="Timing — Day 1: allow roughly 15 minutes meditation plus up to 15 minutes check-in before the slides that follow.",
    )

    slide_sec(
        prs,
        "Day one — morning",
        [
            "Framing, diagnosis and formulation, risk principles, case Leah (shame, family, ND tension, risk).",
            "Psychosynthesis links: parts, meaning alongside risk, brief I–Self and disidentification.",
            "Ends at the morning break before the neurodiversity and Ousmane block.",
        ],
        [
            "Morning material stacks shame, family expectation, and passive risk before the neurodiversity afternoon so trainees feel why formulation cannot be purely categorical.",
            "Leah’s GP letter and voluntary-sector setting mirror common low-intensity entry points where risk is under-reported in referral text.",
        ],
        timing="Timing — Day 1 AM segment: you are heading toward Leah; budget ~40–45 minutes for the full Leah sequence (slides + both task slides + exercise) before the twenty-minute morning break.",
    )

    slide_tc(
        prs,
        "Framing the work",
        [
            "Psychopathology offers shared language across professions and services when that language is used descriptively rather than as identity.",
            "Meaning and risk are held together: one does not cancel assessment, and assessment does not cancel curiosity about what the client cares about.",
            "Uncertainty is ordinary; paralysis and premature closure are both clinical errors in different directions.",
        ],
        [
            "Shared language with GPs and psychiatrists reduces collision between medication, talking therapies, and social care when risk is in play.",
            "Psychosynthesis does not require abandoning DSM or ICD categories; it asks how identification with labels is managed in the relationship.",
        ],
    )

    slide_tc(
        prs,
        "Working with psychopathology",
        [
            "Diagnostic labels describe patterns, not the whole person; clients arrive already shaped by prior assessments and by what families believe those labels mean.",
            "Your formulation integrates symptoms, history, culture, and environment and can change when new information appears.",
            "Thinking clinically means holding categories lightly enough to notice what does not fit, without denying real impairment or risk.",
        ],
        [
            "Collateral from CAMHS, CMHT, or occupational health creates expectation effects; the client may defend or resist a label for good reasons.",
            "Formulation should name maintaining factors outside the skin: housing, work conditions, caring load, discrimination.",
        ],
    )

    slide_tc(
        prs,
        "Risk — core principles",
        [
            "Risk is dynamic: it shifts with isolation, sleep loss, humiliation, substance use, and relationship rupture, not only with expressed emotion.",
            "Intensity of affect is not a one-to-one map to level of danger; quiet compliance can accompany high-lethality planning.",
            "Direct, respectful enquiry about thoughts, intent, plan, means, and protective factors is standard practice, not an escalation in itself.",
            "Thresholds for NHS crisis or emergency services vary by area and gatekeeping; your local pathway still requires a clear clinical rationale in the notes.",
        ],
        [
            "Passive ideation and fleeting intent occupy a wide grey band; documentation should capture verbatim phrases and your clinical reasoning.",
            "Protective factors (children, values, fear of pain) lower but do not remove risk; they belong in formulation and in safety planning.",
            "When self-report shifts between contacts, risk review frequency increases rather than decreases.",
        ],
    )

    slide_to(
        prs,
        'Case Study: "Leah"',
        [
            "Full narrative and tasks: V4-Leah-Shame.docx in the CASE STUDIES folder.",
            "The following slides summarise for teaching; the Word handout is the shared reference for trainees.",
            "Foci: shame, family pressure, tentative neurodiversity themes, shifting risk disclosure, early pacing.",
        ],
        timing="Timing — Leah (Day 1 AM): 40–45 minutes total for all Leah slides including both task slides and paired or small-group work.",
    )

    slide_tc(
        prs,
        "Leah — presentation and context",
        [
            "She has left several retail and hospitality jobs within weeks; busy environments, shift changes, and managerial check-ins overwhelm her.",
            "She withdraws to her room, reverses sleep, drops contact with friends, and describes low energy, self-criticism, and post-social rumination.",
            "A friend raised autism or ADHD; she recognises overload and burnout but refuses assessment, fearing a label would end hope of being “normal.”",
            "Curiosity in the room meets shame and pride; exploration stays descriptive and co-created, not diagnostic.",
        ],
        [
            "Sensory and executive strain can present as “failure to cope” in workplaces that assume uniform tolerance for noise and unpredictability.",
            "Refusing assessment can protect identity while also delaying adjustments that might reduce collapse; both sides belong in formulation.",
        ],
        timing="Timing — Within Leah block: roughly half the Leah window often goes to narrative slides; reserve the remainder for tasks.",
    )

    slide_tc(
        prs,
        "Leah — risk, family pressure, therapy ambivalence",
        [
            "Session two: she denies self-harm; session three after a fight she discloses standing at paracetamol and ibuprofen, wanting everything to stop, then walking away.",
            "She cannot say how long she stood there or whether she handled packets; affect flattens when pressed; she minimises and has not told the GP.",
            "Her mother threatened to stop non-essential spending if she does not find stable work by spring; Leah reads that as conditional love.",
            "Part of her wants honesty and structure; part expects to drop therapy as she dropped jobs; there is no out-of-hours crisis cover in this service.",
            "You must decide in-session whether same-day GP contact is proportionate or whether that risks betrayal and dropout with risk still unassessed elsewhere.",
        ],
        [
            "Shifting self-report across sessions is data: it may reflect shame, cognitive avoidance, or genuine impulse fluctuation.",
            "Family financial threats raise exploitation and emotional safety dynamics alongside suicide risk.",
            "Voluntary-sector contracts often lack medical backup; GP liaison is frequently the available NHS interface.",
        ],
        timing="Timing — Within Leah block: if you are past ~25 minutes from the Leah title slide, move soon to Tasks — Leah.",
    )

    slide_tc(
        prs,
        "Sub-personalities and shame (Leah)",
        [
            "Different parts may hold contradictory positions: one ashamed and hiding risk, one hopeful about therapy, one hostile to parents, one fearful of labels.",
            "Mapping those positions without dramatising can reduce fusion with any single story the client tells in one hour.",
            "Shame often attaches to help-seeking and to risk disclosure; the part that minimises may be trying to preserve dignity.",
        ],
        [
            "Assagioli’s subpersonality language translates to common therapy phrases: inner critic, compliant child, angry adolescent.",
            "Parts work is not a substitute for risk questions; it follows or runs parallel to safeguarding clarity.",
        ],
    )

    slide_tc(
        prs,
        "Meaning alongside risk",
        [
            "Suffering can be meaningful in the client’s life story without minimising danger or delaying safety steps.",
            "Non-identification supports witnessing patterns without the therapist fusing with rescue fantasies or with contempt for the client.",
            "Risk decisions remain ordinary, accountable, and documented; meaning-making does not replace them.",
        ],
        [
            "Firman and Gila’s material on identification helps explain over-responsibility in therapists when clients express passive ideation.",
            "Nelson’s integration theme cautions against romanticising breakdown; the ethical bar stays at safeguarding and consent.",
        ],
    )

    slide_tc(
        prs,
        "I, Self, and disidentification",
        [
            "An observing stance—naming affect and impulse without fusing with them—supports disidentification in session.",
            "Clinical use includes not over-identifying with the client’s narrative or with your own anxiety about getting risk wrong.",
            "Depth and spirit in psychosynthesis still sit inside boundaries, competence, and the duty of care.",
        ],
        [
            "Disidentification is practised in small steps: grounding, sentence stems, timeline work—not instant transpersonal claims.",
            "Therapist countertransference when parents are cast as cruel may obscure the client’s mixed loyalty and dependency.",
        ],
    )

    slide_tc(
        prs,
        "Psychosynthesis perspective (integration)",
        [
            "Distress may reflect fragmentation of subpersonalities; integration is gradual and relationship-dependent.",
            "Containment precedes exploration when shame, risk, or dissociation are live.",
            "Context and inner process are both in the formulation; neither cancels safeguarding or NHS pathways.",
        ],
        [
            "Integration language does not promise linear recovery; it describes direction of travel under supervision.",
            "Firman and Gila’s attention to context supports environmental adjustment as ethical practice, previewing the afternoon block.",
        ],
    )

    slide_tc(
        prs,
        "Tasks — Leah (1)",
        [
            "Formulation under uncertainty: draft a working formulation that holds mood, shame, family context, possible ND-relevant processing differences, and risk—stating explicitly what you do not know and what would change your view.",
            "Immediate clinical decision: what do you do after the cupboard disclosure (same day / before next session)—including whether to involve GP, breach confidentiality to parents, or use another route—and why?",
            TRAINEE_ASSUME,
            INCOMPLETE,
        ],
        [
            "Cupboard disclosure without clear intent or plan still warrants updated risk formulation and often GP awareness when medication access is salient.",
            "Breaching confidentiality to parents risks shame and dropout; withholding may leave domestic risk unmanaged.",
            "Same-day decisions should be recorded with alternatives considered and reasons for the path chosen.",
        ],
        timing="Timing — Leah tasks: allow roughly 20–25 minutes for pair/small-group work on both task slides combined if you are near the 40–45 minute Leah budget.",
    )

    slide_tc(
        prs,
        "Tasks — Leah (2)",
        [
            "Longer-term plan and/or supervision reflection: outline how you would pace the next six to eight sessions if she remains engaged, and name one question you would take to supervision about either risk or family/system pressure.",
            INCOMPLETE,
        ],
        [
            "Pacing balances shame-sensitive exploration with repeated risk review when self-report shifts.",
            "Supervision might examine countertransference toward punitive-sounding parents versus realistic safeguarding concerns.",
        ],
        timing="Timing — End of Leah block: brief plenary or move straight to morning break if time is tight.",
    )

    slide_tc(
        prs,
        "Morning break — 20 minutes",
        [
            "Twenty-minute break before neurodiversity, overload, safeguarding, and Ousmane.",
            "Return time is stated in your cohort schedule; sustained attention drops when breaks are skipped or shortened.",
        ],
        [
            "Cognitive-load theory supports recovery intervals before complex decision tasks such as safeguarding escalation choices.",
        ],
        timing="Timing — Morning break: 20 minutes; confirm return time on the cohort schedule before Ousmane (50–55 min including tasks).",
    )

    slide_sec(
        prs,
        "Day one — afternoon",
        [
            "Neurodiversity, executive load, masking, environment, case Ousmane (overload, children, services).",
            "Practical stabilisation before depth; bifocal vision in tasks.",
            "This block in the file precedes the lunch slide; align to your Saturday clock if needed.",
        ],
        [
            "Afternoon material connects nervous-system load with housing, schooling, and welfare interfaces common in UK practice.",
            "Ousmane’s contradictions train documentation when self-report shifts between contacts.",
        ],
        timing="Timing — Day 1 PM (before lunch in this file): allow ~50–55 minutes for neurodiversity teaching plus full Ousmane sequence including both task slides.",
    )

    slide_tc(
        prs,
        "Neurodiversity — clinical view",
        [
            "Neurodiversity frames difference in processing; impairment is often interaction between person and environment.",
            "Executive functions—planning, inhibition, working memory—explain organisation and timing problems without moral blame.",
            "Sensory overload drives shutdown; masking consumes resources and delays recognition, especially where social expectation is high.",
        ],
        [
            "NICE and NHS services may use diagnostic thresholds for ADHD and autism that lag behind self-identification or community language.",
            "Workplace and education adjustments are legal and clinical levers, not “special treatment” in a pejorative sense.",
        ],
    )

    slide_tc(
        prs,
        "Masking and the cost of compensation",
        [
            "Camouflaging effort in work, parenting, or study leads to exhaustion, late collapse, and misrecognition by brief contacts.",
            "Clients may present as capable in structured appointments while function at home disintegrates.",
            "Formulation should name the price of passing as neurotypical when demand exceeds capacity.",
        ],
        [
            "Burnout after years of masking is often misread as depression without context; history and environment correct that error.",
            "Shame intensifies when failure is attributed to character rather than to load and mismatch.",
        ],
    )

    slide_tc(
        prs,
        "Environment × person",
        [
            "Adjusting context—timing, reminders, sensory load, paperwork support—is part of ethical practice, not a lowering of standards.",
            "When the environment changes, the same person may function markedly differently; formulation updates accordingly.",
            "Liaison with housing, school, or benefits may be as important as insight-oriented work in some weeks.",
        ],
        [
            "Psychosynthesis bifocal vision holds realistic limits and environmental strain alongside meaning and dignity.",
            "Single-parent poverty tightens the margin for error when executive load is high.",
        ],
    )

    slide_tc(
        prs,
        "Clinical challenges (neurodiversity)",
        [
            "Difference is frequently misread as poor motivation or non-compliance in services built on uniform expectations.",
            "Shame accumulates through repeated failure narratives from employers, schools, and prior clinicians.",
            "Your stance needs flexibility without pretending certainty where assessment and support are still pending.",
        ],
        [
            "GPs under time pressure may offer minimal intervention; ADHD waiting lists can exceed twelve months in many areas.",
            "Safeguarding language can be triggered when schools perceive neglect; formulation must separate poverty and overload from intent to harm.",
        ],
    )

    slide_to(
        prs,
        'Case Study: "Ousmane"',
        [
            "Full narrative and tasks: V4-Ousmane-Overload.docx in the CASE STUDIES folder.",
            "Foci: executive overload, parenting, safeguarding edge, contradictory presentation, systems interfaces.",
            "NHS and school pathways constrain what one counselling hour can shift.",
        ],
        timing="Timing — Ousmane (Day 1 PM): 50–55 minutes total for case slides, both task slides, and exercise time.",
    )

    slide_tc(
        prs,
        "Ousmane — overload and contradiction",
        [
            "Executive load: mind goes blank on multi-step tasks, unopened post, phone menus lost once anxiety rises.",
            "Sleep is poor; some days he cannot start meals or wash uniforms despite love for his children and fear of judgement.",
            "Possible ADHD was noted before; assessment never completed after long waits and a dropped triage call.",
            "Between-session texts may sound desperate while sessions sound minimised; channel use may fall outside your service policy.",
        ],
        [
            "Contradiction between capable presentation and functional collapse is common in masking and overload, not proof of manipulation.",
            "Text contact raises boundary and clinical-record questions; agencies differ on whether SMS constitutes part of the record.",
        ],
    )

    slide_tc(
        prs,
        "Ousmane — ethics, services, shifting account",
        [
            "Autonomy and trust versus duty to children: early school involvement may protect or may drive him away from all support.",
            "NHS ADHD waits and unaffordable private assessment limit biological framing even when it would help self-understanding.",
            "GP contact was superficial; mindfulness advice did not address housing or parenting load.",
            "He reports takeaways four nights one week, then “mostly proper meals” the next and cannot align details.",
            "You choose without unanimity or full information: safeguarding conversation, GP re-referral, stabilisation contract, or another proportionate step—and you record the shifting history.",
        ],
        [
            "Children’s services thresholds vary; school concern letters often precede formal section 17 or section 47 processes in England and Wales contexts.",
            "Documentation of contradictory accounts protects you later if risk is retrospectively questioned.",
        ],
    )

    slide_tc(
        prs,
        "Tasks — Ousmane (1)",
        [
            "Formulation under uncertainty: integrate neurodiversity-related and executive-load factors with poverty, parenting load, shame, and service thresholds—stating what remains unknown.",
            "Immediate clinical decision: what is your next move after this session (for example consent for school liaison, GP letter, safeguarding discussion, safety plan only)—with brief rationale and what would make you revise it.",
            TRAINEE_ASSUME,
            INCOMPLETE,
        ],
        [
            "Formulation first clarifies hypotheses about overload versus neglect versus both; without it, liaison can be reactive.",
            "Immediate decisions often hinge on whether children are currently safe enough for brief delay while engagement is secured.",
            "Supervision questions later may target bifocal vision when every systemic door feels closed.",
        ],
        timing="Timing — Ousmane tasks: allow ~25–30 minutes for pair/small-group work across both task slides within the 50–55 minute Ousmane budget.",
    )

    slide_tc(
        prs,
        "Tasks — Ousmane (2)",
        [
            "Longer-term plan and/or supervision reflection: sketch a six-session arc balancing practical stabilisation with emotional work, and name one supervision question on safeguarding boundaries or bifocal vision under extreme strain.",
            INCOMPLETE,
        ],
        [
            "Six-session arcs in overloaded parents often alternate practical wins with shame processing; depth without basics usually fails.",
            "Supervision might examine whether fear of statutory services is distorting your threshold upward or downward.",
        ],
    )

    slide_tc(
        prs,
        "Working practically",
        [
            "Use structure, routine, written prompts, and graded tasks where they reduce executive demand.",
            "Reduce overwhelm before pushing insight; stabilisation supports later meaning-making.",
            "Referral and advocacy—housing, benefits, school—may sit alongside counselling within role boundaries.",
        ],
        [
            "Small repeatable changes track with neuroplasticity themes on day two: repetition shapes habit more than single insights.",
            "Contracts should name how out-of-session contact is used so crisis texts are clinically intentional.",
        ],
        timing="Timing — End of Day 1 PM block in this deck: short consolidation before the 60-minute lunch if your live day still has afternoon teaching after lunch, reorder accordingly.",
    )

    slide_tc(
        prs,
        "Lunch — 1 hour",
        [
            "Lunch is often scheduled 13:00–14:00; confirm the time in your cohort paperwork.",
            "In this deck, day two begins after lunch; your venue may instead continue day one into the late afternoon.",
        ],
        [
            "Post-prandial dips in vigilance are common; risk-heavy material scheduled immediately after lunch may warrant slower pacing.",
        ],
        timing="Timing — Lunch: 60 minutes; confirm start and end times locally (often 13:00–14:00).",
    )

    slide_blank(
        prs,
        [
            "Template blank slide: master/logo styling only, for a short visual pause before day two if you use it live.",
            "Content continues on the next slide; notes carry timing so trainees who receive the deck still have orientation.",
        ],
        timing="Timing — After lunch: optional 30–60 seconds silence or music while people settle before day two opening.",
    )

    slide_tc(
        prs,
        "Day two — opening",
        [
            "Meditation of about fifteen minutes.",
            "Check-in of about fifteen minutes: “What did yesterday bring up for you?”",
            "Then day two morning: trauma, differential with ADHD, dissociation, Emily, neuroplasticity, clinical response.",
        ],
        [
            "Sleep disruption after emotionally loaded training is a common group phenomenon and does not by itself indicate disorder.",
            "Today’s dissociation and escalation teaching assumes yesterday’s shame and risk material is already integrated cognitively.",
        ],
        timing="Timing — Day 2 start: ~15 minutes meditation plus ~15 minutes check-in before trauma and Emily material.",
    )

    slide_sec(
        prs,
        "Day two — morning",
        [
            "Trauma physiology, trauma versus ADHD overlap, dissociation, Emily case, neuroplasticity, phased response.",
            "Ends at afternoon break before Michael and integration themes.",
        ],
        [
            "Dissociation teaching follows trauma framing because defensive absence is easier to understand once nervous-system threat is named.",
            "Emily’s university context adds institutional confidentiality pressure distinct from NHS services.",
        ],
        timing="Timing — Day 2 AM (before afternoon break): allow ~60 minutes for Emily (case + tasks) within a wider morning that also includes trauma overlap teaching, neuroplasticity, and clinical response.",
    )

    slide_tc(
        prs,
        "Trauma — core understanding",
        [
            "Trauma is held in the nervous system as well as in narrative memory; body responses can precede words.",
            "Regulation, perception, and relating change after threat; triggers can flip states quickly.",
            "Clients may oscillate between hyperarousal and shutdown; both carry relational and risk implications.",
        ],
        [
            "PTSD and complex trauma formulations in ICD-11 and DSM-5 still leave grey overlap with anxiety and mood disorders in routine practice.",
            "Therapist pacing regulates arousal in the hour; faster is not always deeper.",
        ],
    )

    slide_tc(
        prs,
        "Trauma presentation and ADHD — overlap",
        [
            "Inattention, restlessness, sleep disruption, and emotional lability appear in both trauma-related states and ADHD.",
            "Premature single explanations waste time; formulation holds both until history, onset, and response to stabilisation clarify the picture.",
            "NHS ADHD assessments are often lengthy; trauma-informed stabilisation can run in parallel without closing either hypothesis.",
        ],
        [
            "Childhood adversity rates are elevated in ADHD cohorts; correlation does not tell you which to treat first in an adult.",
            "Medication decisions belong to psychiatry; counselling formulation still informs what to monitor and report back with consent.",
        ],
    )

    slide_tc(
        prs,
        "Dissociation",
        [
            "Dissociation is often protective after overwhelm; it may show as spacing out, depersonalisation, gaps in time, or inconsistent recall.",
            "Subtle dissociation is easy to miss in brief contacts; risk assessment degrades when continuity of self is unstable.",
            "Strong dissociation correlates with higher risk when intent and behaviour are not fully accessible to conscious report.",
        ],
        [
            "DSM dissociative disorders are under-diagnosed; subthreshold dissociation still matters for safety planning.",
            "Differentiate ordinary distraction from dissociative absence when self-harm or suicide content is present.",
        ],
    )

    slide_tc(
        prs,
        "Containment before exploration",
        [
            "Rapid depth or transpersonal interpretation is unsafe when fragmentation and risk are active.",
            "Stabilisation means relationship reliability, predictable structure, and skills for grounding before trauma narrative work.",
            "Spiritual emergency literature describes crisis states; it is background reading, not a technique checklist for beginners.",
        ],
        [
            "Grof and Clarke volumes inform understanding of psychosis and spirituality interfaces; they do not replace safeguarding law or competence limits.",
            "Nelson’s integration theme warns against forcing insight before the client can stay present.",
        ],
    )

    slide_to(
        prs,
        'Case Study: "Emily"',
        [
            "Full narrative and tasks: V4-Emily-Dissociation.docx in the CASE STUDIES folder.",
            "Foci: dissociation, inconsistent self-harm reporting, university policy tension, escalation under unreliable data.",
            "Do not infer trauma history that is not in the written handout.",
        ],
        timing="Timing — Emily (Day 2 AM): 60 minutes total for case slides, both task slides, and exercise time.",
    )

    slide_tc(
        prs,
        "Emily — mood, dissociation, inconsistent account",
        [
            "Self-harm reportedly “weeks ago” in session one; by session three she cannot tell cutting three days ago from ideation only.",
            "Healed marks on her forearm are shown; she links them to probable dissociation but cannot sequence events.",
            "Detachment episodes last minutes to hours; after one she missed counselling entirely and lost the afternoon.",
            "Sensory disconnection examples include hearing a phone without linking it to herself.",
        ],
        [
            "Inconsistent timelines are clinically meaningful: they may reflect shame, memory fragmentation, or both.",
            "Do not infer a specific trauma history that she has not given; stay with presented material.",
        ],
    )

    slide_tc(
        prs,
        "Emily — risk, services, escalation dilemma",
        [
            "She is unsure she can stay safe; she fears psychiatric assessment, compulsory leave, or being seen as manipulative.",
            "Family values stoicism; disclosure to them is partial. A flat email arrives asking to cancel the next session while saying she is “okay now.”",
            "No GP collateral. Local crisis services often gate on explicit intent or recent attempt; her presentation is partial and fluctuating.",
            "The tutor emailed asking if she is “engaging”; you have not answered. Escalation might save life or destroy alliance; non-escalation might hide risk.",
        ],
        [
            "University counselling sits inside student conduct and safeguarding policies that differ by institution; you need your own service manual.",
            "Flat affect in written communication after dissociation reports is not reassurance; it is ambiguous data.",
        ],
    )

    slide_tc(
        prs,
        "Emily — stabilisation examples (illustrative)",
        [
            "Stabilisation might include brief orientation or grounding in session, a simple sleep or meal routine target, or one negotiated academic boundary such as a supported extension—examples only, not a protocol.",
            "Micro-boundaries on workload reduce overload while trust builds; depth trauma work waits until presence is more reliable.",
            "Risk review frequency increases when self-report is unstable; documentation names what is unknown.",
        ],
        [
            "Stabilisation-first sequencing is consistent with NICE-informed trauma-informed care even when full PTSD treatment waits.",
            "Institutional escalation without consent may be lawful in some serious-risk policies; legality and alliance still conflict clinically.",
        ],
    )

    slide_tc(
        prs,
        "Tasks — Emily (1)",
        [
            "Formulation under uncertainty: link mood, dissociation, academic pressure, family context, and self-harm—explicitly noting contradictions in her account and what you still cannot know.",
            "Immediate clinical decision: given the vague email and unclear recent self-harm, what do you do before the next scheduled contact (outreach, involve others, document and wait)—and what is the ethical trade-off?",
            TRAINEE_ASSUME,
            INCOMPLETE,
        ],
        [
            "Immediate decision-making with partial data mirrors real student services work where crisis lines refuse generic anxiety presentations.",
            "Ethical trade-offs should be written: what you protect (life, autonomy, education) and what you risk (alliance, confidentiality).",
        ],
        timing="Timing — Emily tasks: allow ~28–32 minutes for pair/small-group work across both task slides within the 60-minute Emily budget.",
    )

    slide_tc(
        prs,
        "Tasks — Emily (2)",
        [
            "Longer-term plan and/or supervision reflection: outline stabilisation-first work for the next few sessions if she remains in contact, and one supervision question on escalation thresholds or institutional pressure.",
            INCOMPLETE,
        ],
        [
            "Supervision might test whether fear of being seen as “over-reacting” is lowering your threshold inappropriately.",
            "Stabilisation plans fail if escalation is never reconsidered when new facts appear.",
        ],
        timing="Timing — End of Emily block: brief plenary optional before neuroplasticity and clinical response.",
    )

    slide_tc(
        prs,
        "Neuroplasticity",
        [
            "Repeated experience strengthens neural and emotional patterns—for better and for worse.",
            "Trauma-linked survival responses become automatic; change tends to be gradual and non-linear.",
            "Insight alone rarely shifts embodied habits without repeated relational and behavioural practice.",
        ],
        [
            "This supports realistic homework and between-session structure rather than single breakthrough expectations.",
            "Medication and neurodiversity both modulate plasticity; formulation stays multimodal.",
        ],
        timing="Timing — Day 2 AM: neuroplasticity follows dissociation and Emily work; keep this and the next slide relatively tight so the afternoon break and Michael (60–65 min) stay protected.",
    )

    slide_tc(
        prs,
        "Clinical response",
        [
            "Stabilise before exploratory trauma processing when dissociation or risk is active.",
            "Pace the hour to prevent dysregulation; fewer interpretations when arousal is high.",
            "Risk monitoring continues after clients feel safer because new material often emerges with reduced guarding.",
        ],
        [
            "Collaborative safety planning names triggers, early signs, and help-seeking routes appropriate to your service hours.",
            "Step-down and step-up of frequency should follow risk and functioning, not habit alone.",
        ],
        timing="Timing — Day 2 AM closure: this slide bridges into the twenty-minute afternoon break before Michael and integration.",
    )

    slide_tc(
        prs,
        "Afternoon break — 20 minutes",
        [
            "Twenty-minute break before the afternoon block: Michael case, hold uncertainty while acting, supervision as containment, integration, and closing.",
            "Return time is stated in your cohort schedule.",
        ],
        [
            "The afternoon segment stacks acute-ambiguity risk with documentation and integration; recovery before it supports steadier judgement.",
        ],
        timing="Timing — Afternoon break: 20 minutes; confirm return time before Michael (60–65 min including tasks, hold-uncertainty, supervision, integration).",
    )

    slide_sec(
        prs,
        "Day two — afternoon",
        [
            "Risk under ambiguity (Michael), documentation, supervision, integration, references.",
            "David case removed; Michael carries the acute risk-teaching load for this version.",
        ],
        [
            "Private-sector and stepped-care gaps change what “escalation” means but not documentation standards.",
        ],
        timing="Timing — Day 2 PM: Michael block needs ~60–65 minutes including case slides, tasks, hold uncertainty while acting, supervision, integration, and space before closing.",
    )

    slide_tc(
        prs,
        "Risk decisions in practice",
        [
            "Balance autonomy with duty of care; escalate when risk becomes active against defined criteria.",
            "Supervision supports proportionate decisions; records show questions, answers, and rationale.",
            "Clear documentation protects clients and practitioners when events are later disputed.",
        ],
        [
            "NHS and social care thresholds are not identical to counselling thresholds; know your employer’s policy.",
            "Duty to warn or protect varies by jurisdiction; training curricula should name local law.",
        ],
    )

    slide_to(
        prs,
        'Case Study: "Michael"',
        [
            "Full narrative and tasks: V4-Michael-Risk.docx in the CASE STUDIES folder.",
            "Foci: ambiguous risk below crisis thresholds, contradictory hope cues, private-sector context, hold uncertainty while acting.",
            "Reasonable clinicians may disagree on escalation timing; accountable reasoning is the learning target.",
        ],
        timing="Timing — Michael (Day 2 PM): 60–65 minutes total for case slides, both task slides, and linked discussion.",
    )

    slide_tc(
        prs,
        "Michael — function, contradiction, private context",
        [
            "He missed his last scheduled session, later explaining he had 'lost track of the day,' which fits with his account but adds to the pattern of inconsistency.",
            "Sleep fragments between waking at 3 a.m. and sleeping past midday; eating is patchy; unintended weight loss; hygiene slips on bad weeks.",
            "He works from home on a thin contract; daytime isolation is almost total; no occupational health route.",
            "NHS talking therapies use stepped-care thresholds; private counselling may sit outside those pathways while concern remains real.",
            "Sister offered joint session; he refuses, then says maybe later, and does not follow up—another ambivalent data point.",
        ],
        [
            "A no-show explained as losing the day may reflect cognitive slip in low mood, passive resistance, or both; it widens the file without confirming intent.",
            "Contradictory future orientation versus withdrawal is not intent to die; it is ambiguity that competent clinicians read differently.",
            "Functional markers anchor risk review when words minimise distress.",
        ],
    )

    slide_tc(
        prs,
        "Michael — judgement without certainty",
        [
            "Practitioner unease sits below emergency-service thresholds yet above comfort with simple reassurance.",
            "Autonomy supports continuing exploratory work; duty of care questions whether minimisation is being colluded with.",
            "Insisting on GP contact now might rupture therapy; deferring might leave medical risk unassessed.",
            "Two clinicians could disagree on timing; you still choose an action before the session ends and record unknowns explicitly.",
        ],
        [
            "Documentation separates observed behaviour, client words, and clinician inference—each line carries different evidential weight.",
            "Supervision unpacks whether your anxiety is calibrated to his data or to your own history with loss.",
        ],
    )

    slide_tc(
        prs,
        "Tasks — Michael (1)",
        [
            "Formulation under uncertainty: integrate bereavement, withdrawal, functional change, and the contradictory future hint—stating what supports lower immediate risk and what keeps escalation on the table.",
            "Immediate clinical decision: act now (for example insist on GP contact or boundary shift) or hold and contract—and what would flip your choice overnight?",
            TRAINEE_ASSUME,
            INCOMPLETE,
        ],
        [
            "Overnight flips often involve new intent, substance use, sudden isolation loss, or third-party report of plan.",
            "Private clients still merit NHS medical assessment when physical decline or severe depression is suspected.",
        ],
        timing="Timing — Michael tasks: allow ~30–35 minutes for pair/small-group work across both task slides within the 60–65 minute Michael budget.",
    )

    slide_tc(
        prs,
        "Tasks — Michael (2)",
        [
            "Longer-term plan and/or supervision reflection: outline use of the next four sessions if he continues, and one supervision question separating your anxiety from his data.",
            INCOMPLETE,
        ],
        [
            "Four-session arcs might pair behavioural activation with grief processing if risk stays stable.",
            "Supervision questions on countertransference reduce false certainty in either direction.",
        ],
    )

    slide_tc(
        prs,
        "Hold uncertainty while acting",
        [
            "Uncertainty does not remove duty to decide: you choose outreach level, frequency, documentation, and consultation on incomplete data.",
            "Acting can mean a safety plan, GP encouragement, or family contact with consent—not only emergency services.",
            "Waiting with a clear review date differs from passive drift; both should be explicit in notes.",
        ],
        [
            "Michael’s file illustrates that reasonable disagreement between clinicians is compatible with accountable individual decisions.",
            "Ethical tension stays live; closure narratives are avoided.",
        ],
        timing="Timing — Within Michael block: allow time after tasks for explicit discussion of acting versus waiting with incomplete files.",
    )

    slide_tc(
        prs,
        "Supervision as containment",
        [
            "Supervision holds uncertainty that cannot be fully processed in the hour with the client.",
            "It supports proportionate risk decisions and challenges both rescue and avoidance pulls.",
            "It is a professional requirement in most UK training and practice frameworks, not an optional extra when anxious.",
        ],
        [
            "Group and individual supervision differ in confidentiality; bring anonymised facts consistent with your service policy.",
            "Supervision minutes can be subpoenaed in some proceedings; content should stay factual and proportionate.",
        ],
        timing="Timing — Day 2 PM: supervision material supports the final integration and closing; leave ~15–20 minutes for integration + closing + references if the day ends at 17:00.",
    )

    slide_tc(
        prs,
        "Integration",
        [
            "Hold meaning and risk together across the weekend’s cases without collapsing into either pure pathology or pure optimism.",
            "Recognise limits of the counselling role relative to medicine, statutory services, and employer duties.",
            "Decision-making, documentation, and supervision together form ethical practice—not intuition alone.",
        ],
        [
            "Integration links Leah’s shame, Ousmane’s overload, Emily’s dissociation, and Michael’s ambiguity under one risk–formulation framework.",
            "Psychosynthesis contributions sit alongside safeguarding law and NHS interfaces.",
        ],
        timing="Timing — Pre-closing: integration often runs 10–15 minutes depending on discussion length.",
    )

    slide_tc(
        prs,
        "Closing reflection",
        [
            "A short period of quiet consolidation before leaving can support integration of dense risk and formulation material from the weekend.",
            "Structured peer checkout is optional in many cohorts; psychological safety includes the right not to disclose personal material.",
            "Continued personal activation after risk-focused teaching is a common response; programmes usually signpost pastoral care and supervision rather than large-group therapeutic processing.",
        ],
        [
            "Ambivalence, fatigue, or numbness after risk teaching resemble common stress responses and are not by themselves signs of disorder.",
            "Trainees are encouraged to use programme pastoral routes and their own supervision where personal material has been stirred.",
        ],
        timing="Timing — End of Day 2: allow a few minutes for quiet consolidation and checkout consistent with programme norms.",
    )

    slide_tc(
        prs,
        "Selected references",
        [
            "Nelson, J.E. (1994) Healing the split: Integrating spirit into our understanding of the mentally ill. Albany, NY: State University of New York Press.",
            "Firman, J. and Gila, A. (2002) Psychosynthesis: A psychology of the spirit. Albany, NY: State University of New York Press.",
            "Firman, J. and Gila, A. (1997) The primal wound: A transpersonal view of trauma, addiction, and growth. Albany, NY: State University of New York Press.",
            "Grof, S. and Grof, C. (eds) (1989) Spiritual emergency: When personal transformation becomes a crisis. [Publication details to confirm.]",
            "Clarke, I. (ed.) (2001) Psychosis and spirituality: Exploring the new frontier. [Publication details to confirm.]",
            "Clarke, I. (ed.) (2010) Psychosis and spirituality: Consolidating the new paradigm. [Publication details to confirm.]",
        ],
        [
            "First three references are complete as printed; match your programme’s Harvard punctuation if required.",
            "Grof and Clarke entries lack publisher and place in this master list; confirm in library catalogues before formal essays.",
            "Critical medicalisation literature may be taught separately; only cite with full bibliographic data.",
        ],
        timing="Timing — Final slide: references can be emailed or posted rather than read aloud if time has overrun.",
    )

    # References slide: smaller body
    s = prs.slides[-1]
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 0:
            continue
        if ph.has_text_frame:
            for p in ph.text_frame.paragraphs:
                p.font.size = Pt(REF_BODY_PT)
                for r in p.runs:
                    r.font.size = Pt(REF_BODY_PT)

    prs.save(str(OUT))
    print("Saved", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
