"""Build the Solo Piano musician's guide DOCX from the Markdown source."""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "COMPOSER-OS-SOLO-PIANO-MUSICIANS-GUIDE-V1.0.md"
OUT = ROOT / "Composer-OS-Solo-Piano-Musicians-Guide-V1.0.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
ACCENT = RGBColor(0x0B, 0x3A, 0x5B)


def set_run_font(run, name, size, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_text(paragraph, text, size=11, bold=False, italic=False, color=INK, font="Georgia"):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, font, size, bold, italic, color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, font, size, True, italic, color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", max(9, size - 1), bold, italic, color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, font, size, bold, italic, color)


def style_paragraph(paragraph, before=0, after=8, line=15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line / 12


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_index, row in enumerate(rows):
        for c_index, cell_text in enumerate(row):
            cell = table.rows[r_index].cells[c_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_text(paragraph, cell_text, size=9, bold=(r_index == 0), font="Calibri")
            style_paragraph(paragraph, 1, 1, 12)
    document.add_paragraph()


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Composer OS  ·  Solo Piano Musician's Guide V1.0  ·  Technically verified draft  ·  ")
    set_run_font(run, "Calibri", 8, color=MUTED)
    run = footer.add_run()
    fld_begin = run._r
    # PAGE field
    from docx.oxml import OxmlElement

    def add_field(paragraph):
        run1 = paragraph.add_run()
        r1 = run1._r
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        r1.append(fld1)
        run2 = paragraph.add_run()
        r2 = run2._r
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        r2.append(instr)
        set_run_font(run2, "Calibri", 8, color=MUTED)
        run3 = paragraph.add_run()
        r3 = run3._r
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        r3.append(fld2)

    add_field(footer)

    lines = SRC.read_text().splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(paragraph, line[2:].strip(), size=22, bold=True, color=ACCENT)
            style_paragraph(paragraph, 0, 2, 26)
            index += 1
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph()
            add_text(paragraph, line[3:].strip(), size=16, bold=True, color=ACCENT, font="Calibri")
            style_paragraph(paragraph, 14, 6, 20)
            index += 1
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph()
            add_text(paragraph, line[4:].strip(), size=13, bold=True, color=ACCENT, font="Calibri")
            style_paragraph(paragraph, 10, 4, 16)
            index += 1
            continue
        image = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if image:
            path = ROOT / image.group(2)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(path), width=Inches(6.3))
            style_paragraph(paragraph, 6, 2, 14)
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(caption, image.group(1), size=9, italic=True, color=MUTED, font="Calibri")
            style_paragraph(caption, 0, 8, 12)
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                raw = lines[index].strip().strip("|")
                cells = [cell.strip() for cell in raw.split("|")]
                if not all(set(cell) <= set("-: ") for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                add_table(document, rows)
            continue
        if line.startswith("- "):
            while index < len(lines) and lines[index].startswith("- "):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                add_text(paragraph, "• " + lines[index][2:].strip(), size=11)
                style_paragraph(paragraph, 0, 2, 14)
                index += 1
            continue
        if re.match(r"\d+\. ", line):
            while index < len(lines) and re.match(r"\d+\. ", lines[index]):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                add_text(paragraph, lines[index].strip(), size=11)
                style_paragraph(paragraph, 0, 2, 14)
                index += 1
            continue
        paragraph = document.add_paragraph()
        if line.startswith("**") and line.endswith("**"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(paragraph, line.strip(), size=11)
        style_paragraph(paragraph, 0, 8, 15)
        index += 1

    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
